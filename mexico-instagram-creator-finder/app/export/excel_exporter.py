"""XLSX 导出（openpyxl）。

- 字段宽度合理（根据内容长度自动调整，上限 50，下限 10）
- 长文本自动换行
- URL 字段可点击（hyperlink）
- 冻结首行表头
不导出密码、Cookie 或 Session。
"""

from __future__ import annotations

from pathlib import Path

from app.export.common import records_to_rows
from app.logging_config import get_logger
from app.models import CreatorRecord

logger = get_logger("export.excel")


# URL 字段集合：导出时设为 hyperlink
URL_FIELDS: set[str] = {
    "profile_url",
    "external_url",
    "public_whatsapp_url",
    "linktree_url",
    "beacons_url",
    "profile_pic_url",
}


def export_xlsx(records: list[CreatorRecord], output_path: str | Path) -> Path:
    """
    导出为 XLSX。

    - 字段宽度合理（根据内容长度自动调整，上限 50）
    - 长文本自动换行
    - URL 字段可点击（hyperlink）

    Args:
        records: 已排序的 CreatorRecord 列表
        output_path: 输出文件路径

    Returns:
        实际写入的文件路径。
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    rows, fieldnames = records_to_rows(records)

    wb = Workbook()
    ws = wb.active
    ws.title = "Creators"

    # 表头
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4F81BD")
    for col_idx, field in enumerate(fieldnames, start=1):
        cell = ws.cell(row=1, column=col_idx, value=field)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 数据行
    wrap_alignment = Alignment(wrap_text=True, vertical="top")
    link_font = Font(color="0563C1", underline="single")

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, field in enumerate(fieldnames, start=1):
            value = row.get(field, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = wrap_alignment
            # URL 字段设为 hyperlink
            if field in URL_FIELDS and value and isinstance(value, str) and value.startswith("http"):
                cell.hyperlink = value
                cell.font = link_font

    # 列宽自动调整（10 ~ 50）
    for col_idx, field in enumerate(fieldnames, start=1):
        max_len = len(str(field))
        for row in rows:
            v = row.get(field, "")
            if v is not None:
                # 取首行长度，避免多行文本撑爆列宽
                first_line = str(v).split("\n")[0]
                max_len = max(max_len, len(first_line))
        col_width = min(max(max_len + 2, 10), 50)
        ws.column_dimensions[get_column_letter(col_idx)].width = col_width

    # 冻结首行
    ws.freeze_panes = "A2"

    wb.save(path)
    logger.info("xlsx exported: %s (%d rows)", path, len(rows))
    return path
