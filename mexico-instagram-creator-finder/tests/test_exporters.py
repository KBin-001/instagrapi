"""导出测试。

覆盖 app/export/__init__.py：export_records
- CSV 文件：UTF-8 BOM，中文/西语重音字符不乱码
- JSON 文件：UTF-8，ensure_ascii=False
- XLSX 文件：openpyxl 可读，行数正确，URL 字段为 hyperlink
- 默认排序：total_score 降序 → median_visible_reel_views 降序 → followers 降序
- 不导出密码/Session
- formats 子集
"""

from __future__ import annotations

import csv
import json
from datetime import UTC, datetime
from pathlib import Path

from app.export import export_records
from app.export.common import EXPORT_FIELDS, sort_records
from app.models import (
    AccountTypeClassification,
    ContactInfo,
    CreatorRecord,
    MediaMetrics,
    MexicoSignal,
    NicheClassification,
    ProfileData,
    ScoreResult,
)


def _now() -> datetime:
    return datetime.now(UTC)


def _make_record(
    *,
    username: str,
    follower_count: int,
    total_score: float,
    median_views: float | None,
    full_name: str = "Creator",
    biography: str = "perfume lover",
    level: str = "A",
) -> CreatorRecord:
    """构造完整 CreatorRecord。"""
    profile = ProfileData(
        username=username,
        full_name=full_name,
        biography=biography,
        profile_url=f"https://www.instagram.com/{username}/",
        profile_pic_url="https://example.com/pic.jpg",
        follower_count=follower_count,
        following_count=200,
        media_count=150,
        is_private=False,
        is_verified=False,
        is_business=False,
        category_name="Beauty",
        business_category_name=None,
        external_url=f"https://{username}.mx",
        public_email=f"contact@{username}.mx",
        collected_at=_now(),
    )
    metrics = MediaMetrics(
        recent_media_checked=12,
        recent_reels_checked=5,
        last_post_date=_now(),
        days_since_last_post=3,
        average_likes=500.0,
        median_likes=450.0,
        average_comments=20.0,
        median_comments=18.0,
        average_visible_reel_views=median_views,
        median_visible_reel_views=median_views,
        maximum_visible_reel_views=int(median_views) if median_views else None,
        posting_frequency=3.5,
        reels_view_data_available="available" if median_views else "no_reels",
    )
    contact = ContactInfo(
        public_email=f"contact@{username}.mx",
        public_whatsapp_url=f"https://wa.me/52{username[:9]}",
        external_url=f"https://{username}.mx",
        linktree_url=f"https://linktr.ee/{username}",
        beacons_url=None,
        contact_source="biography_email,biography_wa_me",
        has_public_contact=True,
    )
    niche = NicheClassification(
        primary_niche="perfume",
        niche_scores={"perfume": 0.8},
        niche_signals=["bio keyword"],
        classification_reasons=[],
    )
    mexico = MexicoSignal(
        mexico_confidence_score=0.9,
        mexico_signals=["bio=Mexico"],
        detected_country="Mexico",
        detected_state="Ciudad de México",
        detected_city="polanco",
    )
    account_type = AccountTypeClassification(
        account_type="personal_creator",
        account_type_confidence=0.9,
        account_type_reasons=["default"],
    )
    score = ScoreResult(
        total_score=total_score,
        score_breakdown={"mexico_region": 25.0, "niche_match": 20.0},
        recommendation_level=level,
        recommendation_reasons=["high confidence"],
    )
    return CreatorRecord(
        username=username,
        profile=profile,
        metrics=metrics,
        contact=contact,
        niche=niche,
        mexico=mexico,
        account_type=account_type,
        score=score,
    )


# ---------- CSV 导出 ----------


def test_export_csv_utf8_bom(tmp_path: Path) -> None:
    """CSV 文件：UTF-8 BOM（前 3 字节为 EF BB BF）。"""
    records = [_make_record(username="alice", follower_count=50000, total_score=85.0, median_views=5000)]
    results = export_records(records, tmp_path, formats=["csv"])
    assert "csv" in results
    csv_path = results["csv"]
    # 验证 BOM
    with csv_path.open("rb") as f:
        bom = f.read(3)
    assert bom == b"\xef\xbb\xbf"


def test_export_csv_chinese_and_spanish_not_garbled(tmp_path: Path) -> None:
    """中文/西语重音字符不乱码。"""
    records = [
        _make_record(
            username="alice",
            follower_count=50000,
            total_score=85.0,
            median_views=5000,
            full_name="艾丽斯 García México",
            biography="香水创作者 ñ áéíóú",
        )
    ]
    results = export_records(records, tmp_path, formats=["csv"])
    csv_path = results["csv"]
    # 用 utf-8-sig 读取（自动去除 BOM）
    with csv_path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        rows = list(reader)
        assert len(rows) == 1
        # 验证中文与重音字符原样保留
        assert "艾丽斯" in rows[0]["full_name"]
        assert "García" in rows[0]["full_name"]
        assert "香水创作者" in rows[0]["biography"]
        assert "ñ" in rows[0]["biography"]
        assert "áéíóú" in rows[0]["biography"]


def test_export_csv_contains_all_fields(tmp_path: Path) -> None:
    """CSV 含所有导出字段。"""
    records = [_make_record(username="alice", follower_count=50000, total_score=85.0, median_views=5000)]
    results = export_records(records, tmp_path, formats=["csv"])
    csv_path = results["csv"]
    with csv_path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        headers = reader.fieldnames
        for field in EXPORT_FIELDS:
            assert field in headers, f"CSV 缺少字段: {field}"


# ---------- JSON 导出 ----------


def test_export_json_utf8_no_ascii_escape(tmp_path: Path) -> None:
    """JSON 文件：UTF-8，ensure_ascii=False（含中文/重音字符原样输出）。"""
    records = [
        _make_record(
            username="alice",
            follower_count=50000,
            total_score=85.0,
            median_views=5000,
            full_name="艾丽斯 García",
            biography="香水 México ñ",
        )
    ]
    results = export_records(records, tmp_path, formats=["json"])
    json_path = results["json"]
    # 读取原始字节，验证未转义
    raw = json_path.read_bytes()
    # 中文与重音字符应原样出现（非 \uXXXX 转义）
    assert "艾丽斯".encode() in raw
    assert "García".encode() in raw
    assert "México".encode() in raw
    # 验证可被 json 解析
    data = json.loads(raw.decode("utf-8"))
    assert len(data) == 1
    assert data[0]["full_name"] == "艾丽斯 García"


def test_export_json_structure(tmp_path: Path) -> None:
    """JSON 结构正确。"""
    records = [_make_record(username="alice", follower_count=50000, total_score=85.0, median_views=5000)]
    results = export_records(records, tmp_path, formats=["json"])
    json_path = results["json"]
    data = json.loads(json_path.read_text(encoding="utf-8"))
    assert isinstance(data, list)
    assert len(data) == 1
    assert data[0]["username"] == "alice"
    assert data[0]["total_score"] == 85.0
    assert data[0]["primary_niche"] == "perfume"


# ---------- XLSX 导出 ----------


def test_export_xlsx_readable(tmp_path: Path) -> None:
    """XLSX 文件：openpyxl 可读，行数正确。"""
    records = [
        _make_record(username="alice", follower_count=50000, total_score=85.0, median_views=5000),
        _make_record(username="bob", follower_count=30000, total_score=75.0, median_views=3000),
    ]
    results = export_records(records, tmp_path, formats=["xlsx"])
    xlsx_path = results["xlsx"]
    assert xlsx_path.exists()

    from openpyxl import load_workbook

    wb = load_workbook(xlsx_path)
    ws = wb.active
    # 行数 = 1 表头 + 2 数据
    assert ws.max_row == 3
    wb.close()


def test_export_xlsx_hyperlinks(tmp_path: Path) -> None:
    """XLSX URL 字段为 hyperlink。"""
    records = [_make_record(username="alice", follower_count=50000, total_score=85.0, median_views=5000)]
    results = export_records(records, tmp_path, formats=["xlsx"])
    xlsx_path = results["xlsx"]

    from openpyxl import load_workbook

    wb = load_workbook(xlsx_path)
    ws = wb.active
    # 找到 profile_url 列
    header_row = [cell.value for cell in ws[1]]
    url_col_idx = header_row.index("profile_url") + 1
    # 数据行（第 2 行）
    cell = ws.cell(row=2, column=url_col_idx)
    assert cell.hyperlink is not None
    assert str(cell.hyperlink.target).startswith("https://www.instagram.com/alice/")
    wb.close()


def test_export_xlsx_chinese_not_garbled(tmp_path: Path) -> None:
    """XLSX 中文不乱码。"""
    records = [
        _make_record(
            username="alice",
            follower_count=50000,
            total_score=85.0,
            median_views=5000,
            full_name="艾丽斯",
            biography="香水创作者",
        )
    ]
    results = export_records(records, tmp_path, formats=["xlsx"])
    xlsx_path = results["xlsx"]

    from openpyxl import load_workbook

    wb = load_workbook(xlsx_path)
    ws = wb.active
    header_row = [cell.value for cell in ws[1]]
    name_col_idx = header_row.index("full_name") + 1
    cell = ws.cell(row=2, column=name_col_idx)
    assert cell.value == "艾丽斯"
    wb.close()


# ---------- 默认排序 ----------


def test_default_sort_by_total_score_desc(tmp_path: Path) -> None:
    """默认排序：total_score 降序。"""
    records = [
        _make_record(username="low", follower_count=50000, total_score=50.0, median_views=5000),
        _make_record(username="high", follower_count=50000, total_score=90.0, median_views=5000),
        _make_record(username="mid", follower_count=50000, total_score=70.0, median_views=5000),
    ]
    sorted_records = sort_records(records)
    assert sorted_records[0].username == "high"
    assert sorted_records[1].username == "mid"
    assert sorted_records[2].username == "low"


def test_default_sort_by_median_views_when_score_equal(tmp_path: Path) -> None:
    """total_score 相同时按 median_visible_reel_views 降序。"""
    records = [
        _make_record(username="low_views", follower_count=50000, total_score=80.0, median_views=2000),
        _make_record(username="high_views", follower_count=50000, total_score=80.0, median_views=8000),
    ]
    sorted_records = sort_records(records)
    assert sorted_records[0].username == "high_views"
    assert sorted_records[1].username == "low_views"


def test_default_sort_by_followers_when_score_and_views_equal(tmp_path: Path) -> None:
    """total_score 与 median_views 相同时按 followers 降序。"""
    records = [
        _make_record(username="low_fol", follower_count=30000, total_score=80.0, median_views=5000),
        _make_record(username="high_fol", follower_count=100000, total_score=80.0, median_views=5000),
    ]
    sorted_records = sort_records(records)
    assert sorted_records[0].username == "high_fol"
    assert sorted_records[1].username == "low_fol"


def test_export_csv_applies_default_sort(tmp_path: Path) -> None:
    """导出 CSV 应用默认排序。"""
    records = [
        _make_record(username="low", follower_count=50000, total_score=50.0, median_views=5000),
        _make_record(username="high", follower_count=50000, total_score=90.0, median_views=5000),
    ]
    results = export_records(records, tmp_path, formats=["csv"])
    csv_path = results["csv"]
    with csv_path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        rows = list(reader)
        # 第一行应为 high（分数高）
        assert rows[0]["username"] == "high"
        assert rows[1]["username"] == "low"


# ---------- 不导出密码/Session ----------


def test_no_password_session_in_export_fields() -> None:
    """不导出密码/Session（验证字段列表不含 password/session/cookie）。"""
    forbidden_substrings = ["password", "session", "cookie", "passwd", "pwd", "token", "authorization"]
    for field in EXPORT_FIELDS:
        field_lower = field.lower()
        for sub in forbidden_substrings:
            assert sub not in field_lower, f"导出字段含禁止字段: {field}"


def test_export_csv_no_password_in_content(tmp_path: Path) -> None:
    """CSV 内容不含密码字段。"""
    records = [_make_record(username="alice", follower_count=50000, total_score=85.0, median_views=5000)]
    results = export_records(records, tmp_path, formats=["csv"])
    csv_path = results["csv"]
    content = csv_path.read_text(encoding="utf-8-sig")
    content_lower = content.lower()
    # 不应出现密码相关字段值
    for term in ["password", "ig_password", "sessionid", "cookie"]:
        assert term not in content_lower


# ---------- formats 子集 ----------


def test_export_only_csv(tmp_path: Path) -> None:
    """仅导出 csv 格式。"""
    records = [_make_record(username="alice", follower_count=50000, total_score=85.0, median_views=5000)]
    results = export_records(records, tmp_path, formats=["csv"])
    assert "csv" in results
    assert "json" not in results
    assert "xlsx" not in results
    assert results["csv"].exists()
    # 其他格式文件不应存在
    assert not (tmp_path / "creators.json").exists()
    assert not (tmp_path / "creators.xlsx").exists()


def test_export_only_json(tmp_path: Path) -> None:
    """仅导出 json 格式。"""
    records = [_make_record(username="alice", follower_count=50000, total_score=85.0, median_views=5000)]
    results = export_records(records, tmp_path, formats=["json"])
    assert "json" in results
    assert "csv" not in results
    assert results["json"].exists()


def test_export_all_formats(tmp_path: Path) -> None:
    """默认导出全部三种格式。"""
    records = [_make_record(username="alice", follower_count=50000, total_score=85.0, median_views=5000)]
    results = export_records(records, tmp_path)
    assert "csv" in results
    assert "json" in results
    assert "xlsx" in results
    for path in results.values():
        assert path.exists()


def test_export_empty_records(tmp_path: Path) -> None:
    """空记录列表也能导出（仅表头）。"""
    results = export_records([], tmp_path, formats=["csv", "json"])
    assert results["csv"].exists()
    assert results["json"].exists()
    # CSV 应有表头无数据
    with results["csv"].open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        rows = list(reader)
        assert len(rows) == 0


# ---------- task_id 后缀 ----------


def test_export_with_task_id_suffix(tmp_path: Path) -> None:
    """task_id 用于文件名后缀。"""
    records = [_make_record(username="alice", follower_count=50000, total_score=85.0, median_views=5000)]
    results = export_records(records, tmp_path, formats=["csv"], task_id="abc123")
    assert "csv" in results
    assert "abc123" in results["csv"].name
